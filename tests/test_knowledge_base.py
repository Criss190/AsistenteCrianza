from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from pypdf import PdfWriter

from src.knowledge_base import (
    EXTENSIONES_ADMITIDAS,
    SOLAPAMIENTO_FRAGMENTO,
    TAMANO_FRAGMENTO,
    _leer_texto,
)


class KnowledgeBaseLoaderTests(unittest.TestCase):
    def test_chunking_configuration_is_explicit(self):
        self.assertEqual(TAMANO_FRAGMENTO, 500)
        self.assertEqual(SOLAPAMIENTO_FRAGMENTO, 50)
        self.assertEqual(EXTENSIONES_ADMITIDAS, {".pdf", ".docx", ".txt", ".md"})

    def test_loaders_extract_text_and_location_metadata(self):
        with TemporaryDirectory() as directorio:
            raiz = Path(directorio)
            archivo_txt = raiz / "guia.txt"
            archivo_txt.write_text("Contenido de texto.", encoding="utf-8")

            archivo_md = raiz / "guia.md"
            archivo_md.write_text(
                "# Sección principal\nContenido Markdown.", encoding="utf-8"
            )

            archivo_docx = raiz / "guia.docx"
            documento = Document()
            documento.add_heading("Sección DOCX", level=1)
            documento.add_paragraph("Contenido de Word.")
            documento.save(archivo_docx)

            archivo_pdf = raiz / "guia.pdf"
            escritor = PdfWriter()
            escritor.add_blank_page(width=612, height=792)
            with archivo_pdf.open("wb") as archivo:
                escritor.write(archivo)

            txt = _leer_texto(archivo_txt)
            markdown = _leer_texto(archivo_md)
            docx = _leer_texto(archivo_docx)
            pdf = _leer_texto(archivo_pdf)

        self.assertIn("Contenido de texto.", txt[0][0])
        self.assertEqual(markdown[0][1]["section"], "Sección principal")
        self.assertEqual(docx[0][1]["section"], "Sección DOCX")
        self.assertEqual(pdf[0][1]["page"], 1)


if __name__ == "__main__":
    unittest.main()

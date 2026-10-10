"""Carga, fragmentación, persistencia y recuperación de conocimiento local."""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
from pathlib import Path
import re
import unicodedata
from typing import Any

import chromadb
from chromadb.config import Settings
from docx import Document as WordDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader


EXTENSIONES_ADMITIDAS = {".pdf", ".docx", ".txt", ".md"}
TAMANO_FRAGMENTO = 500
SOLAPAMIENTO_FRAGMENTO = 50
DISTANCIA_MAXIMA = 0.65
NOMBRE_COLECCION = "crianzapp_conocimiento"
PALABRAS_VACIAS = {
    "a", "al", "algo", "cada", "como", "con", "cual", "cuale", "cuando",
    "colombia", "de", "del",
    "donde", "el", "ella", "en", "es", "esta", "este", "hacer", "hay",
    "la", "las", "le", "lo", "los", "me", "mi", "mis", "o", "para",
    "pero", "por", "primer", "que", "se", "son", "su", "sus", "tener",
    "un", "una", "y", "yo", "puedo", "puede", "debo", "deberia",
    "quiero", "ser", "fue", "dos", "tres", "articulo", "evidencia",
    "durante", "intervencione", "actuale",
    "hijo", "hija", "hijos", "hijas", "nino", "nina", "ninos", "ninas",
}
GRUPOS_SINONIMOS = (
    {"pantalla", "television", "tele"},
    {"rabieta", "berrinche", "frustracion"},
    {"emocion", "sentimiento"},
)


def _terminos_relevantes(texto: str) -> set[str]:
    """Obtiene términos de contenido para verificar soporte léxico mínimo."""
    normalizado = unicodedata.normalize("NFD", texto.lower())
    sin_acentos = "".join(
        caracter
        for caracter in normalizado
        if unicodedata.category(caracter) != "Mn"
    )
    terminos = {
        token[:-1] if token.endswith("s") and len(token) > 4 else token
        for token in re.findall(r"[a-z0-9]+", sin_acentos)
    }
    terminos -= PALABRAS_VACIAS
    expandidos = set(terminos)
    for grupo in GRUPOS_SINONIMOS:
        if terminos & grupo:
            expandidos.update(grupo)
    return expandidos


@dataclass(frozen=True)
class FragmentoConocimiento:
    """Fragmento recuperado con su fuente, ubicación y distancia semántica."""

    fuente: str
    texto: str
    chunk_id: str = ""
    page: int | None = None
    section: str | None = None
    distancia: float = 0.0


def _leer_texto(archivo: Path) -> list[tuple[str, dict[str, Any]]]:
    """Extrae texto y metadatos de ubicación de los formatos admitidos."""
    extension = archivo.suffix.lower()
    if extension == ".pdf":
        return [
            (pagina.extract_text() or "", {"page": numero})
            for numero, pagina in enumerate(PdfReader(archivo).pages, start=1)
        ]

    if extension == ".docx":
        documento = WordDocument(archivo)
        secciones: list[tuple[str, dict[str, Any]]] = []
        titulo = archivo.stem
        parrafos: list[str] = []
        for parrafo in documento.paragraphs:
            texto = parrafo.text.strip()
            if not texto:
                continue
            if parrafo.style.name.startswith("Heading"):
                if parrafos:
                    secciones.append(("\n".join(parrafos), {"section": titulo}))
                titulo = texto
                parrafos = [texto]
            else:
                parrafos.append(texto)
        if parrafos:
            secciones.append(("\n".join(parrafos), {"section": titulo}))
        filas_tablas = [
            " | ".join(celda.text.strip() for celda in fila.cells)
            for tabla in documento.tables
            for fila in tabla.rows
        ]
        if filas_tablas:
            secciones.append(
                ("\n".join(filas_tablas), {"section": f"{titulo} (tablas)"})
            )
        return secciones

    texto = archivo.read_text(encoding="utf-8")
    if extension == ".md":
        secciones = []
        encabezado_actual = archivo.stem
        lineas: list[str] = []
        for linea in texto.splitlines():
            encabezado = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", linea)
            if encabezado:
                if lineas:
                    secciones.append(
                        ("\n".join(lineas), {"section": encabezado_actual})
                    )
                encabezado_actual = encabezado.group(1).strip()
                lineas = [linea]
            else:
                lineas.append(linea)
        if lineas:
            secciones.append(("\n".join(lineas), {"section": encabezado_actual}))
        return secciones

    return [(texto, {"section": archivo.stem})]


class BaseConocimientoLocal:
    """Índice Chroma persistente construido desde documentos de conocimiento."""

    def __init__(
        self,
        fragmentos: list[FragmentoConocimiento],
        coleccion: Any,
        client: Any,
    ) -> None:
        self.fragmentos = fragmentos
        self._coleccion = coleccion
        self._client = client

    @classmethod
    def cargar(
        cls,
        directorio: str | Path,
        directorio_vectorstore: str | Path | None = None,
    ) -> "BaseConocimientoLocal":
        """Carga documentos y crea o actualiza el índice vectorial persistente."""
        ruta = Path(directorio).resolve()
        ruta_vectorstore = (
            Path(directorio_vectorstore).resolve()
            if directorio_vectorstore
            else ruta.parent.parent / "vectorstore"
        )
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=TAMANO_FRAGMENTO,
            chunk_overlap=SOLAPAMIENTO_FRAGMENTO,
            add_start_index=True,
        )

        documentos: list[dict[str, Any]] = []
        contenido_hash = hashlib.sha256()
        archivos = sorted(
            archivo
            for archivo in ruta.rglob("*")
            if archivo.is_file() and archivo.suffix.lower() in EXTENSIONES_ADMITIDAS
        )
        for archivo in archivos:
            fuente = archivo.relative_to(ruta).as_posix()
            contenido_hash.update(fuente.encode("utf-8"))
            contenido_hash.update(archivo.read_bytes())
            indice_fragmento = 0
            for texto, ubicacion in _leer_texto(archivo):
                metadatos = {
                    "source": fuente,
                    "filename": archivo.name,
                    **ubicacion,
                }
                for documento in splitter.create_documents(
                    [texto], metadatas=[metadatos]
                ):
                    texto_fragmento = documento.page_content.strip()
                    if texto_fragmento:
                        chunk_id = (
                            f"{hashlib.sha256(fuente.encode()).hexdigest()[:12]}-"
                            f"{indice_fragmento}"
                        )
                        documentos.append(
                            {
                                "text": texto_fragmento,
                                "metadata": {
                                    **documento.metadata,
                                    "chunk_id": chunk_id,
                                },
                            }
                        )
                        indice_fragmento += 1

        huella_corpus = contenido_hash.hexdigest()
        ruta_vectorstore.mkdir(parents=True, exist_ok=True)
        cliente = chromadb.PersistentClient(
            path=str(ruta_vectorstore),
            settings=Settings(anonymized_telemetry=False),
        )
        coleccion = cliente.get_or_create_collection(
            name=NOMBRE_COLECCION,
            metadata={"hnsw:space": "cosine", "corpus_hash": huella_corpus},
        )

        if coleccion.metadata.get("corpus_hash") != huella_corpus or (
            documentos and coleccion.count() == 0
        ):
            cliente.delete_collection(NOMBRE_COLECCION)
            coleccion = cliente.create_collection(
                name=NOMBRE_COLECCION,
                metadata={"hnsw:space": "cosine", "corpus_hash": huella_corpus},
            )
            if documentos:
                coleccion.add(
                    ids=[item["metadata"]["chunk_id"] for item in documentos],
                    documents=[item["text"] for item in documentos],
                    metadatas=[item["metadata"] for item in documentos],
                )

        fragmentos = [
            FragmentoConocimiento(
                fuente=str(item["metadata"]["source"]),
                texto=str(item["text"]),
                chunk_id=str(item["metadata"]["chunk_id"]),
                page=item["metadata"].get("page"),
                section=item["metadata"].get("section"),
            )
            for item in documentos
        ]
        return cls(fragmentos, coleccion, cliente)

    def buscar(self, consulta: str, limite: int = 3) -> list[FragmentoConocimiento]:
        """Recupera los fragmentos más similares y descarta evidencia débil."""
        if not consulta.strip() or limite <= 0 or self._coleccion.count() == 0:
            return []

        resultado = self._coleccion.query(
            query_texts=[consulta],
            n_results=min(max(limite * 4, limite), self._coleccion.count()),
            include=["documents", "metadatas", "distances"],
        )
        documentos = resultado["documents"][0] or []
        metadatos = resultado["metadatas"][0] or []
        distancias = resultado["distances"][0] or []
        terminos_consulta = _terminos_relevantes(consulta)
        if not terminos_consulta:
            return []
        fragmentos: list[FragmentoConocimiento] = []
        for texto, metadata, distancia in zip(documentos, metadatos, distancias):
            if (
                float(distancia) > DISTANCIA_MAXIMA
                or not terminos_consulta & _terminos_relevantes(str(texto))
            ):
                continue
            page = metadata.get("page")
            fragmentos.append(
                FragmentoConocimiento(
                    fuente=str(metadata["source"]),
                    texto=str(texto),
                    chunk_id=str(metadata["chunk_id"]),
                    page=int(page) if page is not None else None,
                    section=metadata.get("section"),
                    distancia=float(distancia),
                )
            )
            if len(fragmentos) == limite:
                break
        return fragmentos
